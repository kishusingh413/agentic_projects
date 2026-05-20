from __future__ import annotations

import csv
import json
import os
from typing import Iterable, Mapping, Sequence

from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, Field, ValidationError, field_validator

from services.errors import (
    ErrorContext,
    ProductsFileFormatError,
    ProductsFileValidationError,
)


class ProductRecord(BaseModel):
    """Validated representation of a single product row/object."""

    sku: str = Field(min_length=1)

    # Allow additional product attributes; we will embed them into prompts.
    model_config = ConfigDict(extra="allow")

    @field_validator("sku")
    @classmethod
    def normalize_sku(cls, v: str) -> str:
        cleaned = v.strip()
        if not cleaned:
            raise ValueError("sku must be non-empty")
        return cleaned


def _infer_format_from_path(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".csv":
        return "csv"
    if ext in (".xls", ".xlsx"):
        return "excel"
    if ext == ".json":
        return "json"
    raise ProductsFileFormatError(
        f"Unsupported file extension: {ext or '(none)'}",
        context=ErrorContext(action="infer_format", payload={"path": path}),
    )


def _find_field_name(headers: Sequence[str], candidates: Sequence[str]) -> str | None:
    normalized = {h.strip().lower(): h for h in headers}
    for c in candidates:
        key = c.strip().lower()
        if key in normalized:
            return normalized[key]
    return None


def _to_product_prompt(record: Mapping[str, object]) -> str:
    """Convert a validated record into a single string for task prompts."""
    sku_value = str(record.get("sku", "")).strip()
    if not sku_value:
        # Should not happen if validation is correct, but keep a guard.
        return "Unnamed SKU"

    # Embed other attributes to give the LLM structured context.
    # Keep it stable and concise: omit empty/None, sort keys.
    extra_items: list[str] = []
    for key in sorted(record.keys()):
        if key == "sku":
            continue
        value = record.get(key)
        if value is None:
            continue
        if isinstance(value, str) and not value.strip():
            continue
        extra_items.append(f"{key}={value}")

    if not extra_items:
        return sku_value

    return f"{sku_value} | " + " | ".join(extra_items)


def _validate_products(records: Iterable[Mapping[str, object]], *, action: str) -> list[str]:
    out: list[str] = []
    for idx, raw in enumerate(records):
        try:
            validated = ProductRecord.model_validate(raw)
        except ValidationError as e:
            raise ProductsFileValidationError(
                "Product row validation failed.",
                context=ErrorContext(
                    action=action,
                    payload={"row_index": idx, "errors": e.errors(), "raw": dict(raw)},
                ),
            ) from e
        out.append(_to_product_prompt(validated.model_dump()))
    return out


def load_products_from_json(path: str) -> list[str]:
    with open(path, "r", encoding="utf-8-sig") as f:
        data = json.load(f)

    if isinstance(data, dict):
        # Common shapes: { "products": [...] } or { "items": [...] }
        for key in ("products", "items", "data"):
            if key in data and isinstance(data[key], list):
                data = data[key]
                break

    if not isinstance(data, list):
        raise ProductsFileValidationError(
            "JSON must be a list of product objects (or an object containing a products array).",
            context=ErrorContext(action="load_json", payload={"path": path, "type": type(data).__name__}),
        )

    def normalize_record(item: Mapping[str, object]) -> Mapping[str, object]:
        # Support common SKU key alternatives.
        sku_value = None
        if "sku" in item:
            sku_value = item.get("sku")
        elif "product" in item:
            sku_value = item.get("product")
        elif "product_id" in item:
            sku_value = item.get("product_id")
        elif "id" in item:
            sku_value = item.get("id")

        if sku_value is None:
            # Let schema validation fail with a clear error.
            return dict(item)

        normalized: dict[str, object] = dict(item)
        normalized["sku"] = sku_value
        return normalized

    records = (
        normalize_record(item)
        for item in data
        if isinstance(item, dict)
    )
    return _validate_products(records, action="load_json")


def load_products_from_csv(path: str, *, sku_column: str | None = None) -> list[str]:
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        if reader.fieldnames is None:
            raise ProductsFileValidationError(
                "CSV is missing headers.",
                context=ErrorContext(action="load_csv", payload={"path": path}),
            )

        fieldnames = [str(h) for h in reader.fieldnames]

        if sku_column is not None:
            if sku_column not in fieldnames:
                raise ProductsFileValidationError(
                    f"CSV sku_column '{sku_column}' not found in headers.",
                    context=ErrorContext(
                        action="load_csv",
                        payload={"path": path, "sku_column": sku_column, "headers": fieldnames},
                    ),
                )
            sku_field = sku_column
        else:
            sku_field = _find_field_name(fieldnames, ["sku", "product", "product_id", "id"])
            if sku_field is None:
                raise ProductsFileValidationError(
                    "CSV must include a SKU column (e.g., 'sku' or 'product_id').",
                    context=ErrorContext(action="load_csv", payload={"path": path, "headers": fieldnames}),
                )

        # Normalize keys: map sku column to 'sku'.
        def records() -> Iterable[Mapping[str, object]]:
            for row in reader:
                if not row:
                    continue
                normalized: dict[str, object] = {}
                for k, v in row.items():
                    if k is None:
                        continue
                    key = str(k).strip()
                    if not key:
                        continue
                    if key == sku_field:
                        normalized["sku"] = v
                        continue
                    normalized[key] = v
                yield normalized

        return _validate_products(records(), action="load_csv")


def load_products_from_excel(
    path: str,
    *,
    sku_column: str | None = None,
) -> list[str]:
    wb = load_workbook(path, read_only=True, data_only=True)
    sheet = wb.active

    # Read header row.
    rows = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration as e:
        raise ProductsFileValidationError(
            "Excel sheet is empty.",
            context=ErrorContext(action="load_excel", payload={"path": path}),
        ) from e

    headers: list[str] = [str(h).strip() for h in header_row if h is not None and str(h).strip()]
    if not headers:
        raise ProductsFileValidationError(
            "Excel is missing headers.",
            context=ErrorContext(action="load_excel", payload={"path": path}),
        )

    if sku_column is not None:
        if sku_column not in headers:
            raise ProductsFileValidationError(
                f"Excel sku_column '{sku_column}' not found in headers.",
                context=ErrorContext(
                    action="load_excel",
                    payload={"path": path, "sku_column": sku_column, "headers": headers},
                ),
            )
        sku_field = sku_column
    else:
        found = _find_field_name(headers, ["sku", "product", "product_id", "id"])
        if found is None:
            raise ProductsFileValidationError(
                "Excel must include a SKU column (e.g., 'sku' or 'product_id').",
                context=ErrorContext(action="load_excel", payload={"path": path, "headers": headers}),
            )
        sku_field = found

    def records() -> Iterable[Mapping[str, object]]:
        for row in rows:
            # Skip blank rows (all None).
            if row is None:
                continue
            values = list(row)
            if all(v is None or (isinstance(v, str) and not v.strip()) for v in values):
                continue

            normalized: dict[str, object] = {}
            for i, header in enumerate(headers):
                value = values[i] if i < len(values) else None
                if header == sku_field:
                    normalized["sku"] = value
                    continue
                normalized[header] = value
            yield normalized

    return _validate_products(records(), action="load_excel")


def load_products_from_file(
    path: str,
    *,
    file_format: str | None = None,
    sku_column: str | None = None,
) -> list[str]:
    if not path.strip():
        raise ProductsFileValidationError(
            "products-file path is empty.",
            context=ErrorContext(action="load_products_from_file", payload={"path": path}),
        )

    if not os.path.exists(path):
        raise ProductsFileValidationError(
            "products-file does not exist.",
            context=ErrorContext(action="load_products_from_file", payload={"path": path}),
        )

    format_name = file_format.strip().lower() if file_format else _infer_format_from_path(path)
    if format_name == "csv":
        return load_products_from_csv(path, sku_column=sku_column)
    if format_name == "excel":
        return load_products_from_excel(path, sku_column=sku_column)
    if format_name == "json":
        return load_products_from_json(path)

    raise ProductsFileFormatError(
        "Unsupported products file format.",
        context=ErrorContext(action="load_products_from_file", payload={"path": path, "format": format_name}),
    )

